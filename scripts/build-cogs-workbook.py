"""왕징양다리양꼬치 매출원가계산서 엑셀 생성 스크립트.

실행: python scripts/build-cogs-workbook.py [출력경로]
기본 출력: docs/operations/왕징_매출원가계산서.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/operations/왕징_매출원가계산서.xlsx"

# ---------------------------------------------------------------- 스타일
FONT = "맑은 고딕"
DARK, CORAL, TINT, SAND = "1E1928", "F54B1E", "FEEAE2", "F5EBE1"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color=DARK)
F_SECTION = Font(name=FONT, size=12, bold=True, color=CORAL)
F_NOTE = Font(name=FONT, size=9, color="716969")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_HEAD = PatternFill("solid", fgColor=DARK)
FILL_SUB = PatternFill("solid", fgColor=SAND)
FILL_TOTAL = PatternFill("solid", fgColor=TINT)
THIN = Side(style="thin", color="D0C8C0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0#;(#,##0.0#);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"

RED_FILL = PatternFill("solid", fgColor="F8D7D5")
ORANGE_FILL = PatternFill("solid", fgColor="FDEBD0")
GREEN_FILL = PatternFill("solid", fgColor="D5EFE3")


def style(c, font=F_BASE, fmt=None, fill=None, align=None, border=True):
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = BOX
    return c


def header_row(ws, row, headers, col=1, height=32):
    for i, h in enumerate(headers):
        style(ws.cell(row, col + i, h), F_HEAD, fill=FILL_HEAD, align=CENTER)
    ws.row_dimensions[row].height = height


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None, span="A1:H1"):
    ws[span.split(":")[0]] = text
    ws[span.split(":")[0]].font = F_TITLE
    ws.row_dimensions[1].height = 28
    if sub:
        c = ws.cell(2, ws[span.split(":")[0]].column, sub)
        c.font = F_NOTE


def add_list_dv(ws, formula, rng, strict=True, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = strict
    if not strict:
        dv.errorStyle = "information"
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "입력 안내", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def status_cf(ws, rng):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("위험",{first}))'], fill=RED_FILL, font=Font(name=FONT, color="C8322A", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("주의",{first}))'], fill=ORANGE_FILL, font=Font(name=FONT, color="9C5700", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("적정",{first}))'], fill=GREEN_FILL, font=Font(name=FONT, color="187A56", bold=True)))


# ---------------------------------------------------------------- 기준 데이터
STORES = ["판교점", "모란점"]
CATEGORIES = [
    ("양고기", "식재료", "양다리·양갈비·양꼬치(완제품 꼬치 포함)·볶음용 양고기"),
    ("기타육류", "식재료", "돼지고기(꿔바로우·육사용), 닭고기 등"),
    ("수산물", "식재료", "대하·새우꼬치·민물가재(소룽샤)"),
    ("채소·과일", "식재료", "가지·토마토·감자·피망·고수·레몬 등"),
    ("두부·계란·곡류·면", "식재료", "두부·건두부·계란·쌀·면류·분모자"),
    ("냉동·가공식품", "식재료", "물만두·꽃빵·피쉬볼 등 완제품"),
    ("양념·향신료·소스", "식재료", "쯔란·마라소스·두반장·춘장·간장·설탕 등"),
    ("유지·기타식재료", "식재료", "식용유·참기름·땅콩·기본찬 재료"),
    ("주류", "주류·음료", "백주(연태구냥·설원 등)·맥주·하이볼용 위스키"),
    ("음료·얼음", "주류·음료", "탄산수·콜라·사이다·얼음"),
    ("숯·소모품·포장", "소모품", "숯·꼬치·물티슈·앞치마·포장용기"),
]
NCAT = len(CATEGORIES)
CAT_FIRST, CAT_LAST = 4, 4 + NCAT - 1  # 기본설정 시트 분류표 행 범위

# (코드, 분류, 품목명, 규격/원산지, 주거래처, 구매단위, 구매가격(공급가), 구매단위당 용량, 기준단위, 수율, 과세구분)
ITEMS = [
    ("M01", "양고기", "양다리(뼈포함·냉동)", "호주산", "양고기 거래처", "kg", 14000, 1000, "g", 0.95, "면세"),
    ("M02", "양고기", "양갈비 프렌치랙(냉동)", "호주산", "양고기 거래처", "kg", 32000, 1000, "g", 0.92, "면세"),
    ("M03", "양고기", "생양꼬치(완제품)", "100개입", "양꼬치 거래처", "박스", 45000, 100, "개", 1, "면세"),
    ("M04", "양고기", "양념양꼬치(완제품)", "100개입", "양꼬치 거래처", "박스", 52000, 100, "개", 1, "과세"),
    ("M05", "양고기", "양갈비살꼬치(완제품)", "100개입", "양꼬치 거래처", "박스", 62000, 100, "개", 1, "면세"),
    ("M06", "양고기", "양어깨살(볶음·탕용)", "호주산", "양고기 거래처", "kg", 18000, 1000, "g", 0.9, "면세"),
    ("P01", "기타육류", "돼지 등심(꿔바로우용)", "국내산", "정육 거래처", "kg", 11000, 1000, "g", 0.9, "면세"),
    ("P02", "기타육류", "돼지 앞다리 채(육사용)", "국내산", "정육 거래처", "kg", 9000, 1000, "g", 0.95, "면세"),
    ("S01", "수산물", "새우꼬치(완제품)", "50개입", "양꼬치 거래처", "박스", 35000, 50, "개", 1, "면세"),
    ("S02", "수산물", "대하(냉동)", "수입", "수산 거래처", "kg", 24000, 1000, "g", 0.85, "면세"),
    ("S03", "수산물", "민물가재(냉동)", "중국산", "수산 거래처", "kg", 18000, 1000, "g", 0.95, "면세"),
    ("V01", "채소·과일", "가지", "국내산", "채소 거래처", "kg", 4500, 1000, "g", 0.9, "면세"),
    ("V02", "채소·과일", "토마토", "국내산", "채소 거래처", "kg", 5000, 1000, "g", 0.95, "면세"),
    ("V03", "채소·과일", "감자", "국내산", "채소 거래처", "kg", 2500, 1000, "g", 0.85, "면세"),
    ("V04", "채소·과일", "피망", "국내산", "채소 거래처", "kg", 6000, 1000, "g", 0.85, "면세"),
    ("V05", "채소·과일", "오이", "국내산", "채소 거래처", "kg", 3500, 1000, "g", 0.95, "면세"),
    ("V06", "채소·과일", "양파", "국내산", "채소 거래처", "kg", 1800, 1000, "g", 0.9, "면세"),
    ("V07", "채소·과일", "대파", "국내산", "채소 거래처", "kg", 3500, 1000, "g", 0.85, "면세"),
    ("V08", "채소·과일", "깐마늘", "국내산", "채소 거래처", "kg", 9000, 1000, "g", 1, "면세"),
    ("V09", "채소·과일", "생강", "국내산", "채소 거래처", "kg", 8000, 1000, "g", 0.85, "면세"),
    ("V10", "채소·과일", "고수", "국내산", "채소 거래처", "kg", 12000, 1000, "g", 0.8, "면세"),
    ("V11", "채소·과일", "청경채", "국내산", "채소 거래처", "kg", 5000, 1000, "g", 0.85, "면세"),
    ("V12", "채소·과일", "숙주", "국내산", "채소 거래처", "kg", 2500, 1000, "g", 0.95, "면세"),
    ("V13", "채소·과일", "배추", "국내산", "채소 거래처", "kg", 1500, 1000, "g", 0.8, "면세"),
    ("V14", "채소·과일", "레몬", "수입", "채소 거래처", "개", 600, 1, "개", 1, "면세"),
    ("T01", "두부·계란·곡류·면", "두부", "국내산", "식자재마트", "kg", 3500, 1000, "g", 1, "면세"),
    ("T02", "두부·계란·곡류·면", "건두부(포두부)", "중국산", "중국식자재", "kg", 9000, 1000, "g", 1, "과세"),
    ("T03", "두부·계란·곡류·면", "계란", "30구", "식자재마트", "판", 8000, 30, "개", 1, "면세"),
    ("T04", "두부·계란·곡류·면", "쌀", "20kg", "식자재마트", "포대", 60000, 20000, "g", 1, "면세"),
    ("T05", "두부·계란·곡류·면", "냉면 사리", "2kg", "식자재마트", "봉", 6000, 2000, "g", 1, "과세"),
    ("T06", "두부·계란·곡류·면", "옥수수면", "1kg", "중국식자재", "봉", 5000, 1000, "g", 1, "과세"),
    ("T07", "두부·계란·곡류·면", "분모자·당면", "1kg", "중국식자재", "봉", 6000, 1000, "g", 1, "과세"),
    ("F01", "냉동·가공식품", "물만두(냉동)", "1kg", "중국식자재", "봉", 7000, 1000, "g", 1, "과세"),
    ("F02", "냉동·가공식품", "꽃빵(냉동)", "30개입", "중국식자재", "봉", 7500, 30, "개", 1, "과세"),
    ("F03", "냉동·가공식품", "피쉬볼·어묵", "1kg", "중국식자재", "봉", 8000, 1000, "g", 1, "과세"),
    ("J01", "양념·향신료·소스", "쯔란 시즈닝", "1kg", "중국식자재", "봉", 15000, 1000, "g", 1, "과세"),
    ("J02", "양념·향신료·소스", "마라탕 소스", "500g", "중국식자재", "봉", 6000, 500, "g", 1, "과세"),
    ("J03", "양념·향신료·소스", "마라샹궈 소스", "500g", "중국식자재", "봉", 7000, 500, "g", 1, "과세"),
    ("J04", "양념·향신료·소스", "두반장", "1kg", "중국식자재", "통", 8000, 1000, "g", 1, "과세"),
    ("J05", "양념·향신료·소스", "첨면장(춘장)", "1kg", "중국식자재", "통", 6000, 1000, "g", 1, "과세"),
    ("J06", "양념·향신료·소스", "간장", "1.8L", "식자재마트", "병", 6000, 1800, "ml", 1, "과세"),
    ("J07", "양념·향신료·소스", "식초", "1.8L", "식자재마트", "병", 3000, 1800, "ml", 1, "과세"),
    ("J08", "양념·향신료·소스", "설탕", "3kg", "식자재마트", "봉", 5000, 3000, "g", 1, "과세"),
    ("J09", "양념·향신료·소스", "고춧가루", "1kg", "식자재마트", "봉", 18000, 1000, "g", 1, "면세"),
    ("J10", "양념·향신료·소스", "건고추", "1kg", "중국식자재", "봉", 16000, 1000, "g", 1, "면세"),
    ("J11", "양념·향신료·소스", "화자오(산초)", "1kg", "중국식자재", "봉", 30000, 1000, "g", 1, "과세"),
    ("J12", "양념·향신료·소스", "감자전분", "1kg", "식자재마트", "봉", 3500, 1000, "g", 1, "과세"),
    ("O01", "유지·기타식재료", "식용유", "18L", "식자재마트", "말", 55000, 18000, "ml", 1, "과세"),
    ("O02", "유지·기타식재료", "참기름", "1.8L", "식자재마트", "병", 25000, 1800, "ml", 1, "과세"),
    ("O03", "유지·기타식재료", "땅콩(기본찬)", "1kg", "식자재마트", "봉", 10000, 1000, "g", 1, "면세"),
    ("L01", "주류", "연태구냥 500ml", "34도", "주류 도매상", "병", 12000, 1, "개", 1, "과세"),
    ("L02", "주류", "연태구냥 250ml", "34도", "주류 도매상", "병", 6500, 1, "개", 1, "과세"),
    ("L03", "주류", "연태구냥 125ml", "34도", "주류 도매상", "병", 3800, 1, "개", 1, "과세"),
    ("L04", "주류", "설원 450ml", "30도", "주류 도매상", "병", 8000, 1, "개", 1, "과세"),
    ("L05", "주류", "설원 250ml", "30도", "주류 도매상", "병", 4800, 1, "개", 1, "과세"),
    ("L06", "주류", "공부가주 500ml", "33도", "주류 도매상", "병", 17000, 1, "개", 1, "과세"),
    ("L07", "주류", "노주탄 500ml", "33도", "주류 도매상", "병", 9500, 1, "개", 1, "과세"),
    ("L08", "주류", "칭다오 640ml", "4.7도", "주류 도매상", "병", 2600, 1, "개", 1, "과세"),
    ("L09", "주류", "하얼빈 500ml", "4.3도", "주류 도매상", "병", 2000, 1, "개", 1, "과세"),
    ("L10", "주류", "산토리 가쿠빈 700ml", "하이볼용", "주류 도매상", "병", 26000, 700, "ml", 1, "과세"),
    ("L11", "주류", "제임슨 700ml", "하이볼용", "주류 도매상", "병", 28000, 700, "ml", 1, "과세"),
    ("L12", "주류", "짐빔 700ml", "하이볼용", "주류 도매상", "병", 22000, 700, "ml", 1, "과세"),
    ("L13", "주류", "커티삭 700ml", "하이볼용", "주류 도매상", "병", 16000, 700, "ml", 1, "과세"),
    ("L14", "주류", "봄베이 사파이어 750ml", "하이볼용", "주류 도매상", "병", 30000, 750, "ml", 1, "과세"),
    ("L15", "주류", "연태구냥 500ml(하이볼용)", "ml 단위 원가용", "주류 도매상", "병", 12000, 500, "ml", 1, "과세"),
    ("B01", "음료·얼음", "탄산수 1.5L", "하이볼용", "식자재마트", "병", 1500, 1500, "ml", 1, "과세"),
    ("B02", "음료·얼음", "콜라 355ml", "캔", "식자재마트", "캔", 900, 1, "개", 1, "과세"),
    ("B03", "음료·얼음", "사이다 355ml", "캔", "식자재마트", "캔", 900, 1, "개", 1, "과세"),
    ("B04", "음료·얼음", "얼음", "3kg", "식자재마트", "봉", 3000, 3000, "g", 1, "과세"),
    ("C01", "숯·소모품·포장", "숯(야자숯)", "10kg", "숯 거래처", "박스", 25000, 10000, "g", 1, "과세"),
    ("C02", "숯·소모품·포장", "물티슈", "1000개", "소모품 거래처", "박스", 20000, 1000, "개", 1, "과세"),
    ("C03", "숯·소모품·포장", "일회용 앞치마", "100개", "소모품 거래처", "박스", 8000, 100, "개", 1, "과세"),
    ("C04", "숯·소모품·포장", "포장용기 세트", "100개", "소모품 거래처", "박스", 25000, 100, "개", 1, "과세"),
]

# (메뉴분류, 메뉴명, 원가구분, 판매가) — 홈페이지 메뉴판(data/menu.csv) 기준
MENUS = [
    ("시그니처 양다리", "양다리구이 (대)", "식재료", 90000),
    ("시그니처 양다리", "양다리구이 (중)", "식재료", 80000),
    ("양꼬치·세트", "고급양갈비", "식재료", 30000),
    ("양꼬치·세트", "생양꼬치", "식재료", 17000),
    ("양꼬치·세트", "양념양꼬치", "식재료", 18000),
    ("양꼬치·세트", "양갈비살꼬치", "식재료", 18000),
    ("양꼬치·세트", "새우꼬치", "식재료", 18000),
    ("중국요리·탕", "꿔바로우", "식재료", 20000),
    ("중국요리·탕", "가지튀김", "식재료", 18000),
    ("중국요리·탕", "토마토계란볶음", "식재료", 16000),
    ("중국요리·탕", "마파두부", "식재료", 15000),
    ("중국요리·탕", "향라대하", "식재료", 22000),
    ("중국요리·탕", "어향육사", "식재료", 19000),
    ("중국요리·탕", "경장육사", "식재료", 19000),
    ("중국요리·탕", "마라탕", "식재료", 18000),
    ("중국요리·탕", "소룽샤", "식재료", 38000),
    ("중국요리·탕", "지삼선", "식재료", 18000),
    ("중국요리·탕", "건두부볶음", "식재료", 16000),
    ("중국요리·탕", "오이무침", "식재료", 12000),
    ("중국요리·탕", "쯔란양고기", "식재료", 28000),
    ("중국요리·탕", "마라샹궈", "식재료", 32000),
    ("중국요리·탕", "건두부무침", "식재료", 16000),
    ("중국요리·탕", "양탕", "식재료", 15000),
    ("식사·면·디저트", "계란볶음밥", "식재료", 8000),
    ("식사·면·디저트", "가지볶음밥", "식재료", 8000),
    ("식사·면·디저트", "옥수수온면", "식재료", 8000),
    ("식사·면·디저트", "냉면", "식재료", 8000),
    ("식사·면·디저트", "물만두", "식재료", 8000),
    ("식사·면·디저트", "꽃빵튀김", "식재료", 8000),
    ("주류·하이볼", "연태구냥 500ml", "주류·음료", 40000),
    ("주류·하이볼", "연태구냥 250ml", "주류·음료", 20000),
    ("주류·하이볼", "연태구냥 125ml", "주류·음료", 12000),
    ("주류·하이볼", "설원 450ml", "주류·음료", 25000),
    ("주류·하이볼", "설원 250ml", "주류·음료", 15000),
    ("주류·하이볼", "공부가주 500ml", "주류·음료", 50000),
    ("주류·하이볼", "노주탄 500ml", "주류·음료", 30000),
    ("주류·하이볼", "칭다오 맥주 640ml", "주류·음료", 7000),
    ("주류·하이볼", "하얼빈 맥주 500ml", "주류·음료", 7000),
    ("주류·하이볼", "산토리하이볼", "주류·음료", 8000),
    ("주류·하이볼", "제임슨하이볼", "주류·음료", 8000),
    ("주류·하이볼", "짐빔하이볼", "주류·음료", 7000),
    ("주류·하이볼", "커티삭하이볼", "주류·음료", 6000),
    ("주류·하이볼", "봄베이하이볼", "주류·음료", 8000),
    ("주류·하이볼", "연태하이볼", "주류·음료", 7000),
]


def highball(spirit, ml=45):
    return [(spirit, ml), ("탄산수 1.5L", 150), ("얼음", 150), ("레몬", 0.125)]


# 메뉴명 → [(품목명, 1인분 사용량(기준단위))] — 표준 레시피 예시(추정치)
RECIPES = {
    "양다리구이 (대)": [("양다리(뼈포함·냉동)", 1800), ("쯔란 시즈닝", 30), ("양파", 100), ("고수", 20)],
    "양다리구이 (중)": [("양다리(뼈포함·냉동)", 1400), ("쯔란 시즈닝", 25), ("양파", 80), ("고수", 15)],
    "고급양갈비": [("양갈비 프렌치랙(냉동)", 300), ("쯔란 시즈닝", 10)],
    "생양꼬치": [("생양꼬치(완제품)", 10), ("쯔란 시즈닝", 10)],
    "양념양꼬치": [("양념양꼬치(완제품)", 10), ("쯔란 시즈닝", 8)],
    "양갈비살꼬치": [("양갈비살꼬치(완제품)", 10), ("쯔란 시즈닝", 10)],
    "새우꼬치": [("새우꼬치(완제품)", 10)],
    "꿔바로우": [("돼지 등심(꿔바로우용)", 250), ("감자전분", 120), ("식용유", 150), ("설탕", 60), ("식초", 50)],
    "가지튀김": [("가지", 400), ("감자전분", 100), ("식용유", 150), ("설탕", 30), ("간장", 20)],
    "토마토계란볶음": [("토마토", 350), ("계란", 4), ("식용유", 40), ("설탕", 15), ("대파", 20)],
    "마파두부": [("두부", 400), ("돼지 앞다리 채(육사용)", 80), ("두반장", 40), ("화자오(산초)", 3), ("식용유", 40), ("대파", 20), ("감자전분", 15)],
    "향라대하": [("대하(냉동)", 250), ("건고추", 20), ("화자오(산초)", 5), ("깐마늘", 30), ("식용유", 80), ("대파", 30)],
    "어향육사": [("돼지 앞다리 채(육사용)", 250), ("피망", 80), ("양파", 80), ("두반장", 30), ("설탕", 20), ("식초", 20), ("식용유", 60), ("감자전분", 20)],
    "경장육사": [("돼지 앞다리 채(육사용)", 250), ("첨면장(춘장)", 50), ("대파", 120), ("건두부(포두부)", 80), ("식용유", 50), ("설탕", 15)],
    "마라탕": [("마라탕 소스", 120), ("양어깨살(볶음·탕용)", 100), ("청경채", 80), ("숙주", 80), ("배추", 100), ("분모자·당면", 100), ("피쉬볼·어묵", 100), ("두부", 100)],
    "소룽샤": [("민물가재(냉동)", 600), ("마라샹궈 소스", 100), ("건고추", 20), ("깐마늘", 40), ("식용유", 100), ("대파", 30)],
    "지삼선": [("가지", 250), ("감자", 200), ("피망", 100), ("감자전분", 50), ("식용유", 150), ("간장", 25), ("설탕", 15), ("깐마늘", 15)],
    "건두부볶음": [("건두부(포두부)", 250), ("피망", 60), ("양파", 60), ("돼지 앞다리 채(육사용)", 60), ("식용유", 40), ("간장", 20)],
    "오이무침": [("오이", 350), ("깐마늘", 20), ("식초", 30), ("참기름", 10), ("설탕", 10), ("고춧가루", 5)],
    "쯔란양고기": [("양어깨살(볶음·탕용)", 250), ("쯔란 시즈닝", 15), ("양파", 100), ("고수", 15), ("건고추", 10), ("식용유", 50)],
    "마라샹궈": [("마라샹궈 소스", 100), ("양어깨살(볶음·탕용)", 150), ("대하(냉동)", 100), ("감자", 100), ("청경채", 80), ("건두부(포두부)", 80), ("분모자·당면", 80), ("식용유", 80)],
    "건두부무침": [("건두부(포두부)", 250), ("오이", 80), ("고수", 15), ("식초", 20), ("참기름", 10), ("고춧가루", 5)],
    "양탕": [("양어깨살(볶음·탕용)", 180), ("대파", 30), ("고수", 10), ("생강", 5)],
    "계란볶음밥": [("쌀", 200), ("계란", 2), ("대파", 15), ("식용유", 20)],
    "가지볶음밥": [("쌀", 200), ("가지", 80), ("계란", 1), ("식용유", 25), ("간장", 10)],
    "옥수수온면": [("옥수수면", 200), ("청경채", 30), ("대파", 10), ("간장", 15)],
    "냉면": [("냉면 사리", 200), ("오이", 30), ("계란", 0.5), ("식초", 15), ("설탕", 15)],
    "물만두": [("물만두(냉동)", 300), ("간장", 10), ("식초", 10)],
    "꽃빵튀김": [("꽃빵(냉동)", 8), ("식용유", 100), ("설탕", 10)],
    "연태구냥 500ml": [("연태구냥 500ml", 1)],
    "연태구냥 250ml": [("연태구냥 250ml", 1)],
    "연태구냥 125ml": [("연태구냥 125ml", 1)],
    "설원 450ml": [("설원 450ml", 1)],
    "설원 250ml": [("설원 250ml", 1)],
    "공부가주 500ml": [("공부가주 500ml", 1)],
    "노주탄 500ml": [("노주탄 500ml", 1)],
    "칭다오 맥주 640ml": [("칭다오 640ml", 1)],
    "하얼빈 맥주 500ml": [("하얼빈 500ml", 1)],
    "산토리하이볼": highball("산토리 가쿠빈 700ml"),
    "제임슨하이볼": highball("제임슨 700ml"),
    "짐빔하이볼": highball("짐빔 700ml"),
    "커티삭하이볼": highball("커티삭 700ml"),
    "봄베이하이볼": highball("봄베이 사파이어 750ml"),
    "연태하이볼": highball("연태구냥 500ml(하이볼용)", 30),
}

_item_names = {i[2] for i in ITEMS}
_cat_names = {c[0] for c in CATEGORIES}
for i in ITEMS:
    assert i[1] in _cat_names, i
for m, rows in RECIPES.items():
    assert m in {x[1] for x in MENUS}, m
    for name, _ in rows:
        assert name in _item_names, (m, name)

# 시트별 용량(행)
ITEM_FIRST, ITEM_LAST = 4, 203
RCP_FIRST, RCP_LAST = 4, 603
MENU_FIRST, MENU_LAST = 6, 65
SALES_FIRST, SALES_LAST = 4, 403
PUR_FIRST, PUR_LAST = 4, 1503
INV_FIRST, INV_LAST = 4, 403

S = "'기본설정'"
YEAR, MONTH, STORE = f"{S}!$C$4", f"{S}!$C$5", f"{S}!$C$6"
ALLOW, TGT_FOOD, TGT_DRINK, TOL, DEDUCT, GAP = (f"{S}!$C${r}" for r in (7, 8, 9, 10, 12, 13))
CAT_RNG, GRP_RNG = f"{S}!$G${CAT_FIRST}:$G${CAT_LAST}", f"{S}!$H${CAT_FIRST}:$H${CAT_LAST}"
STORE_LIST = f"={S}!$E$5:$E$10"


def rng(sheet, col, first, last):
    return f"'{sheet}'!${col}${first}:${col}${last}"


def item_lookup(col, key):
    return f"INDEX({rng('식자재단가', col, ITEM_FIRST, ITEM_LAST)},MATCH({key},{rng('식자재단가', 'C', ITEM_FIRST, ITEM_LAST)},0))"


def by_store(sum_args, store_col_rng):
    """조회지점이 '전체'면 지점 조건 없이, 아니면 지점 조건을 붙여 SUMIFS."""
    return f'IF({STORE}="전체",SUMIFS({sum_args}),SUMIFS({sum_args},{store_col_rng},{STORE}))'


wb = Workbook()

# ================================================================ 1. 사용법
ws = wb.active
ws.title = "사용법"
widths(ws, {"A": 3, "B": 24, "C": 90})
title(ws, "왕징양다리양꼬치 매출원가계산서 — 사용법", span="B1:C1")
ws["B2"] = "한 파일로 1년(1~12월)을 관리합니다. 매일 매출·매입을 쌓고, 매월 말 재고만 세면 원가율이 자동 계산됩니다."
ws["B2"].font = F_NOTE
guide = [
    ("■ 시트 구성", None),
    ("기본설정", "상호·연도·조회월·조회지점, 목표 원가율, 분류표, 의제매입 공제율. 매월 [조회월]만 바꾸면 계산서가 그 달로 바뀝니다."),
    ("식자재단가", "품목 마스터. 거래명세서 단가가 바뀌면 [구매가격]만 수정 → 레시피·메뉴원가·재고금액이 모두 자동 갱신됩니다."),
    ("레시피", "메뉴 1인분에 들어가는 재료와 사용량(g·ml·개). 품목은 드롭다운으로 선택하세요."),
    ("메뉴원가표", "메뉴별 1인분 원가·원가율·공헌이익, 월 판매수량 입력 시 이론원가와 메뉴분석(스타/정리검토)까지 표시."),
    ("매출일보", "매일 마감 후 1줄(지점별). 카드·현금·이체·배달 매출, 그중 주류·음료 매출, 객수 입력."),
    ("매입장", "거래명세서·영수증 받을 때마다 1줄. 품목을 고르면 분류·단위·과세구분이 자동으로 채워집니다."),
    ("재고실사", "매월 말일 영업 종료 후 남은 수량을 해당 월 칸에 입력(구매단위 기준: kg, 박스, 병 등)."),
    ("매출원가계산서", "조회월·조회지점의 매출 → 매출원가(분류별) → 원가율 판정 → 매출총이익 → 이론 vs 실제 → 의제매입 참고."),
    ("연간요약", "1~12월 원가율 추이를 한 표로. 재고실사를 안 한 달은 '미입력'으로 표시됩니다."),
    ("", None),
    ("■ 셀 색상 규칙", None),
    ("노란 바탕 + 파란 글씨", "직접 입력하는 칸입니다. 여기만 입력하세요."),
    ("흰 바탕 + 검은 글씨", "자동 계산(수식) 칸입니다. 지우거나 덮어쓰지 마세요."),
    ("초록 글씨", "다른 시트에서 자동으로 불러온 값입니다."),
    ("", None),
    ("■ 운영 루틴", None),
    ("매일 (3분)", "① 매출일보에 POS 마감 금액 입력  ② 들어온 거래명세서를 매입장에 입력"),
    ("매주 (10분)", "식자재단가 시트에서 단가가 오른 품목 확인(변동률 10% 이상은 빨간색). 양고기·주류 단가는 특히 자주 확인."),
    ("매월 말일 (30분)", "① 재고실사 해당 월 칸 입력  ② 기본설정 조회월 변경  ③ 메뉴원가표에 POS 메뉴별 판매수량 입력  ④ 매출원가계산서 확인"),
    ("", None),
    ("■ 계산 원리", None),
    ("매출원가", "기초재고 + 당월매입 − 기말재고  (전월 말 재고 = 이번 달 기초재고, 자동 연결)"),
    ("원가율", "매출원가 ÷ 매출(공급가액, VAT 제외).  식재료는 음식매출과, 주류·음료는 주류매출과 비교합니다."),
    ("이론원가 vs 실제원가", "이론원가 = 레시피 원가 × 판매수량. 실제원가가 이론보다 5% 이상 크면 로스·과다투입·서비스·도난·매입 오기를 점검하세요."),
    ("금액 기준", "모든 원가는 공급가액(부가세 제외)으로 계산합니다. 매출도 ÷1.1 해서 공급가액으로 비교합니다."),
    ("", None),
    ("■ 왕징 맞춤 관리 포인트", None),
    ("양고기", "원가의 가장 큰 비중. 양다리는 해동 손실(수율 95%)을 반영했고, 수율은 실제로 무게를 달아 고치면 더 정확해집니다."),
    ("완제품 꼬치", "양꼬치는 '박스(100개)' 단위로 매입하고 레시피는 '개' 단위로 사용합니다. 1인분 꼬치 개수가 다르면 레시피에서 수정."),
    ("주류", "주류는 원가율이 낮은 효자 상품입니다. 음식과 분리해서 보는 이유입니다. 하이볼은 위스키 45ml(연태 30ml) 기준."),
    ("숯·소모품", "숯은 메뉴별로 나누기 어려워 '소모품' 원가로 따로 집계하고, 계산서에서 포함/제외 원가율을 둘 다 보여줍니다."),
    ("의제매입세액", "양고기·채소·두부·계란 등 면세 농축수산물 매입은 부가세 신고 때 공제받을 수 있습니다(계산서 6번 참고)."),
    ("", None),
    ("■ 미리 넣어둔 값에 대해", None),
    ("단가·레시피", "식자재 단가와 레시피 사용량은 2026년 도매 시세 기준 '예시 추정치'입니다. 반드시 실제 거래명세서와 주방 계량값으로 바꿔주세요."),
    ("메뉴·판매가", "홈페이지 메뉴판(44개 메뉴) 가격을 그대로 넣었습니다."),
    ("예시 행", "매출일보·매입장의 '예시(삭제)' 행은 입력 형식 참고용입니다. 확인 후 행을 지우고 사용하세요."),
    ("지점", "판교점·모란점이 기본입니다. 기본설정 지점목록에서 이름을 바꾸거나 추가할 수 있습니다."),
]
r = 4
for k, v in guide:
    if v is None:
        ws.cell(r, 2, k).font = F_SECTION
    else:
        ws.cell(r, 2, k).font = F_BOLD
        c = ws.cell(r, 3, v)
        c.font = F_BASE
        c.alignment = LEFT
    r += 1
for rr, (fill, font) in {17: (FILL_INPUT, F_INPUT), 18: (None, F_BASE), 19: (None, F_LINK)}.items():
    c = ws.cell(rr, 2)
    if fill:
        c.fill = fill
    c.font = Font(name=FONT, size=10, bold=True, color=font.color)

# ================================================================ 2. 기본설정
ws = wb.create_sheet("기본설정")
widths(ws, {"A": 2, "B": 30, "C": 22, "D": 2, "E": 14, "F": 2, "G": 20, "H": 12, "I": 46, "J": 2, "K": 26, "L": 12, "M": 2, "N": 10, "O": 10, "P": 14, "Q": 12})
title(ws, "기본설정", "노란 칸만 수정하세요. 매월 [조회월]을 바꾸면 매출원가계산서가 해당 월로 바뀝니다.", span="B1:I1")
settings = [
    (3, "상호", "왕징양다리양꼬치", None, None),
    (4, "연도", 2026, "0", None),
    (5, "조회월 (1~12)", 10, '0"월"', None),
    (6, "조회지점", "전체", None, None),
    (7, "잡재료 가산율 (양념·기본찬)", 0.05, PCT, "레시피에 넣지 않은 소량 양념·기본찬(땅콩·짜사이·양배추)·쯔란 리필 등을 음식 원가에 일괄 가산하는 비율. 일반적으로 3~7%."),
    (8, "목표 식재료 원가율", 0.35, PCT, "양꼬치·중식 전문점 통상 식재료 원가율 30~38% 수준(업계 통상치). 사장님 목표로 수정하세요."),
    (9, "목표 주류·음료 원가율", 0.30, PCT, "백주·맥주 판매가 대비 매입가 비율. 통상 25~35%."),
    (10, "경고 허용폭 (%p)", 0.03, PCT, "실제 원가율이 목표보다 높으면 '주의', 목표+허용폭보다 높으면 '위험'."),
    (11, "사업자 유형 (의제매입)", "개인(과세표준 2억 이하)", None, None),
    (12, "의제매입세액 공제율", f"=INDEX($L$4:$L$6,MATCH(C11,$K$4:$K$6,0))", "0.00%", None),
    (13, "이론-실제 허용 차이율", 0.05, PCT, "실제원가가 레시피 기준 이론원가보다 이 비율 이상 크면 로스 점검 신호."),
]
for row, label, val, fmt, note in settings:
    style(ws.cell(row, 2, label), F_BOLD, fill=FILL_SUB, align=LEFT)
    c = ws.cell(row, 3, val)
    is_formula = isinstance(val, str) and val.startswith("=")
    style(c, F_BASE if is_formula else F_INPUT, fmt, None if is_formula else FILL_INPUT, CENTER)
    if note:
        c.comment = Comment(note, "왕징 원가표")
ws["B15"] = "※ 목표 원가율은 업계 통상치를 참고한 출발값입니다. 3개월 실적을 본 뒤 사장님 기준으로 조정하세요."
ws["B15"].font = F_NOTE
ws["B16"] = "※ 의제매입세액 공제율은 2026년 현행 기준(개인 음식점 과표 2억 이하 9/109 특례). 세법 개정 시 K~L열을 수정하세요."
ws["B16"].font = F_NOTE

style(ws["E3"], F_HEAD, fill=FILL_HEAD, align=CENTER).value = "지점목록"
style(ws["E4"], F_BASE, align=CENTER).value = "전체"
for i in range(6):
    c = ws.cell(5 + i, 5, STORES[i] if i < len(STORES) else None)
    style(c, F_INPUT, fill=FILL_INPUT, align=CENTER)

header_row(ws, 3, ["분류", "원가그룹", "포함 품목 예시"], col=7)
for i, (cat, grp, desc) in enumerate(CATEGORIES):
    style(ws.cell(CAT_FIRST + i, 7, cat), F_INPUT, fill=FILL_INPUT, align=LEFT)
    style(ws.cell(CAT_FIRST + i, 8, grp), F_INPUT, fill=FILL_INPUT, align=CENTER)
    style(ws.cell(CAT_FIRST + i, 9, desc), F_BASE, align=LEFT)

header_row(ws, 3, ["사업자 유형", "공제율"], col=11)
for i, (k, v) in enumerate([("개인(과세표준 2억 이하)", "=9/109"), ("개인(과세표준 2억 초과)", "=8/108"), ("법인", "=6/106")]):
    style(ws.cell(4 + i, 11, k), F_BASE, align=LEFT)
    style(ws.cell(4 + i, 12, v), F_INPUT, "0.00%", FILL_INPUT, CENTER)

for col, head, vals in [(14, "과세구분", ["과세", "면세"]), (15, "결제방법", ["카드", "계좌이체", "현금", "외상"]),
                        (16, "증빙", ["세금계산서", "계산서(면세)", "카드전표", "현금영수증", "간이영수증"]),
                        (17, "원가그룹", ["식재료", "주류·음료", "소모품"])]:
    style(ws.cell(3, col, head), F_HEAD, fill=FILL_HEAD, align=CENTER)
    for i, v in enumerate(vals):
        style(ws.cell(4 + i, col, v), F_BASE, align=CENTER)

dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12")
dv.error, dv.errorTitle = "1~12 사이 숫자를 입력하세요.", "조회월"
ws.add_data_validation(dv)
dv.add("C5")
add_list_dv(ws, "=$E$4:$E$10", "C6")
add_list_dv(ws, "=$K$4:$K$6", "C11")
add_list_dv(ws, "=$Q$4:$Q$6", f"H{CAT_FIRST}:H{CAT_LAST}")
ws.freeze_panes = "A3"

# ================================================================ 3. 식자재단가
ws = wb.create_sheet("식자재단가")
title(ws, "식자재 단가표 (품목 마스터)", "구매가격은 공급가액(부가세 제외). 단가가 바뀌면 [이전 구매가]에 옛 가격을 옮기고 [구매가격]을 고치면 변동률이 표시됩니다. ※ 미리 넣은 단가는 예시 추정치 — 실제 거래명세서 단가로 교체하세요.")
heads = ["코드", "분류", "품목명", "규격·원산지", "주거래처", "구매단위", "구매가격\n(공급가, 원)", "구매단위당\n용량", "기준\n단위", "기준단위당\n단가(원)", "수율\n(손질·해동후)", "실사용 단가\n(원/기준단위)", "과세\n구분", "단가\n변경일", "이전\n구매가", "단가\n변동률"]
header_row(ws, 3, heads, height=40)
widths(ws, dict(zip("ABCDEFGHIJKLMNOP", [7, 18, 26, 14, 14, 9, 13, 11, 7, 12, 11, 13, 7, 11, 11, 9])))
input_cols = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 13, 14, 15}
fmts = {7: NUM, 8: "#,##0", 10: NUM1, 11: "0%", 12: NUM1, 14: DATE, 15: NUM, 16: PCT}
for r in range(ITEM_FIRST, ITEM_LAST + 1):
    item = ITEMS[r - ITEM_FIRST] if r - ITEM_FIRST < len(ITEMS) else None
    vals = {}
    if item:
        code, cat, name, spec, vendor, unit, price, qty, base, yld, tax = item
        vals = {1: code, 2: cat, 3: name, 4: spec, 5: vendor, 6: unit, 7: price, 8: qty, 9: base, 11: yld, 13: tax}
    vals[10] = f'=IF(OR(G{r}="",H{r}="",H{r}=0),"",G{r}/H{r})'
    vals[12] = f'=IF(J{r}="","",J{r}/IF(OR(K{r}="",K{r}=0),1,K{r}))'
    vals[16] = f'=IF(OR(O{r}="",O{r}=0,G{r}=""),"",G{r}/O{r}-1)'
    for col in range(1, 17):
        c = ws.cell(r, col, vals.get(col))
        inp = col in input_cols
        style(c, F_INPUT if inp else F_BASE, fmts.get(col), FILL_INPUT if inp else None, CENTER if col not in (3, 4, 5) else LEFT)
ws.cell(ITEM_FIRST, 11).comment = Comment("수율 = 실제 쓰는 양 ÷ 구매한 양.\n예) 양다리 1kg 해동 후 950g → 95%.\n뼈·껍질·손질 손실이 클수록 실사용 단가가 올라갑니다.", "왕징 원가표")
add_list_dv(ws, f"={CAT_RNG}", f"B{ITEM_FIRST}:B{ITEM_LAST}")
add_list_dv(ws, '"g,ml,개"', f"I{ITEM_FIRST}:I{ITEM_LAST}")
add_list_dv(ws, f"={S}!$N$4:$N$5", f"M{ITEM_FIRST}:M{ITEM_LAST}")
ws.conditional_formatting.add(f"P{ITEM_FIRST}:P{ITEM_LAST}", FormulaRule(formula=[f'AND(ISNUMBER(P{ITEM_FIRST}),P{ITEM_FIRST}>=0.1)'], font=Font(name=FONT, color="C8322A", bold=True)))
ws.conditional_formatting.add(f"P{ITEM_FIRST}:P{ITEM_LAST}", FormulaRule(formula=[f'AND(ISNUMBER(P{ITEM_FIRST}),P{ITEM_FIRST}<=-0.1)'], font=Font(name=FONT, color="187A56", bold=True)))
ws.freeze_panes = "D4"
ws.auto_filter.ref = f"A3:P{ITEM_LAST}"

# ================================================================ 4. 레시피
ws = wb.create_sheet("레시피")
title(ws, "표준 레시피 (메뉴 1인분 기준)", "메뉴명·재료·사용량만 입력하면 단위·단가·원가가 자동 계산됩니다. 사용량 단위는 식자재단가의 기준단위(g·ml·개)와 같아야 합니다. ※ 사용량은 예시 추정치 — 실제 계량값으로 수정하세요.")
header_row(ws, 3, ["메뉴명", "재료(품목명)", "1인분\n사용량", "단위", "실사용 단가\n(원/단위)", "재료원가\n(원)", "메모"], height=40)
widths(ws, {"A": 22, "B": 26, "C": 10, "D": 7, "E": 13, "F": 12, "G": 30})
recipe_rows = [(m, it, q) for (_, m, _, _) in MENUS for (it, q) in RECIPES.get(m, [])]
for r in range(RCP_FIRST, RCP_LAST + 1):
    rec = recipe_rows[r - RCP_FIRST] if r - RCP_FIRST < len(recipe_rows) else (None, None, None)
    for col, v in ((1, rec[0]), (2, rec[1]), (3, rec[2]), (7, None)):
        style(ws.cell(r, col, v), F_INPUT, NUM1 if col == 3 else None, FILL_INPUT, LEFT if col != 3 else CENTER)
    style(ws.cell(r, 4, f'=IF(B{r}="","",IFERROR({item_lookup("I", f"B{r}")},"?"))'), F_LINK, align=CENTER)
    style(ws.cell(r, 5, f'=IF(B{r}="","",IFERROR({item_lookup("L", f"B{r}")},"품목없음"))'), F_LINK, NUM1, align=CENTER)
    style(ws.cell(r, 6, f'=IF(OR(B{r}="",C{r}=""),"",IF(ISNUMBER(E{r}),C{r}*E{r},0))'), F_BASE, NUM)
add_list_dv(ws, f"='메뉴원가표'!$B${MENU_FIRST}:$B${MENU_LAST}", f"A{RCP_FIRST}:A{RCP_LAST}", strict=False)
add_list_dv(ws, f"={rng('식자재단가', 'C', ITEM_FIRST, ITEM_LAST)}", f"B{RCP_FIRST}:B{RCP_LAST}")
ws.conditional_formatting.add(f"E{RCP_FIRST}:E{RCP_LAST}", CellIsRule(operator="equal", formula=['"품목없음"'], fill=RED_FILL))
ws.freeze_panes = "B4"
ws.auto_filter.ref = f"A3:G{RCP_LAST}"

# ================================================================ 5. 메뉴원가표
ws = wb.create_sheet("메뉴원가표")
title(ws, "메뉴별 원가표 · 메뉴 분석", "판매가는 VAT 포함 메뉴판 가격. 원가율은 공급가(판매가÷1.1) 기준. [월 판매수량]은 POS '메뉴별 판매현황'에서 조회월 수량을 입력하세요.", span="A1:Q1")
# 상단 요약: 라벨(3행) / 값(4행)
summary = [("B", "총 판매수량", f"=SUM(M{MENU_FIRST}:M{MENU_LAST})", "#,##0"),
           ("D", "가중평균 공헌이익(원)", f"=IFERROR(SUMPRODUCT(L{MENU_FIRST}:L{MENU_LAST},M{MENU_FIRST}:M{MENU_LAST})/B4,0)", NUM),
           ("F", "인기 기준수량(70% 법칙)", f"=IFERROR(B4/COUNTA(B{MENU_FIRST}:B{MENU_LAST})*0.7,0)", "#,##0.0"),
           ("H", "이론원가 합계(원)", f"=SUM(O{MENU_FIRST}:O{MENU_LAST})", NUM),
           ("J", "이론 매출(공급가)", f"=SUM(N{MENU_FIRST}:N{MENU_LAST})", NUM),
           ("L", "이론원가율", "=IFERROR(H4/J4,0)", PCT)]
for col, label, f, fmt in summary:
    style(ws[f"{col}3"], F_BOLD, fill=FILL_SUB, align=CENTER).value = label
    style(ws[f"{col}4"], F_BOLD, fmt, FILL_TOTAL, CENTER).value = f
ws.row_dimensions[3].height = 28
heads = ["메뉴분류", "메뉴명", "원가구분", "판매가\n(VAT포함)", "판매가\n(공급가)", "주재료원가\n(레시피)", "잡재료\n가산", "1인분\n원가", "원가율", "목표\n원가율", "판정", "공헌이익\n(1인분)", "월\n판매수량", "매출\n(공급가)", "이론원가", "매출\n비중", "메뉴분석"]
header_row(ws, 5, heads, height=40)
widths(ws, dict(zip("ABCDEFGHIJKLMNOPQ", [15, 20, 10, 11, 11, 11, 9, 10, 8, 8, 13, 11, 9, 12, 11, 8, 28])))
RCP_A, RCP_F = rng("레시피", "A", RCP_FIRST, RCP_LAST), rng("레시피", "F", RCP_FIRST, RCP_LAST)
for r in range(MENU_FIRST, MENU_LAST + 1):
    m = MENUS[r - MENU_FIRST] if r - MENU_FIRST < len(MENUS) else (None, None, None, None)
    for col, v, fmt in ((1, m[0], None), (2, m[1], None), (3, m[2], None), (4, m[3], NUM), (13, None, "#,##0")):
        style(ws.cell(r, col, v), F_INPUT, fmt, FILL_INPUT, CENTER if col > 2 else LEFT)
    f = {
        5: (f'=IF(D{r}="","",D{r}/1.1)', NUM),
        6: (f'=IF(B{r}="","",SUMIF({RCP_A},B{r},{RCP_F}))', NUM),
        7: (f'=IF(B{r}="","",IF(C{r}="식재료",F{r}*{ALLOW},0))', NUM),
        8: (f'=IF(B{r}="","",F{r}+G{r})', NUM),
        9: (f'=IF(OR(B{r}="",D{r}=""),"",IFERROR(H{r}/E{r},0))', PCT),
        10: (f'=IF(B{r}="","",IF(C{r}="주류·음료",{TGT_DRINK},{TGT_FOOD}))', PCT),
        11: (f'=IF(B{r}="","",IF(D{r}="","판매가 미입력",IF(F{r}=0,"레시피 미입력",IF(I{r}>J{r}+{TOL},"위험",IF(I{r}>J{r},"주의","적정")))))', None),
        12: (f'=IF(OR(B{r}="",D{r}=""),"",E{r}-H{r})', NUM),
        14: (f'=IF(OR(M{r}="",E{r}=""),"",E{r}*M{r})', NUM),
        15: (f'=IF(OR(M{r}="",H{r}=""),"",H{r}*M{r})', NUM),
        16: (f'=IF(N{r}="","",IFERROR(N{r}/$J$4,0))', PCT),
        17: (f'=IF(OR(M{r}="",M{r}=0,L{r}=""),"",IF(AND(M{r}>=$F$4,L{r}>=$D$4),"스타 (유지·강조)",IF(M{r}>=$F$4,"인기·저마진 (원가↓·가격↑ 검토)",IF(L{r}>=$D$4,"고마진·비인기 (홍보 강화)","저인기·저마진 (정리 검토)"))))', None),
    }
    for col, (formula, fmt) in f.items():
        style(ws.cell(r, col, formula), F_BASE, fmt, align=CENTER if col in (9, 10, 11, 16) else (LEFT if col == 17 else None))
add_list_dv(ws, '"식재료,주류·음료"', f"C{MENU_FIRST}:C{MENU_LAST}")
status_cf(ws, f"K{MENU_FIRST}:K{MENU_LAST}")
ws.conditional_formatting.add(f"K{MENU_FIRST}:K{MENU_LAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("미입력",K{MENU_FIRST}))'], font=Font(name=FONT, color="716969", italic=True)))
ws.conditional_formatting.add(f"Q{MENU_FIRST}:Q{MENU_LAST}", FormulaRule(formula=[f'LEFT(Q{MENU_FIRST},2)="스타"'], font=Font(name=FONT, color="187A56", bold=True)))
ws.conditional_formatting.add(f"Q{MENU_FIRST}:Q{MENU_LAST}", FormulaRule(formula=[f'LEFT(Q{MENU_FIRST},3)="저인기"'], font=Font(name=FONT, color="C8322A", bold=True)))
ws.cell(MENU_LAST + 2, 1, "※ 메뉴분석: 판매수량이 '인기 기준수량' 이상이면 인기, 공헌이익이 '가중평균 공헌이익' 이상이면 고마진 (메뉴 엔지니어링 기법). 음식·주류를 함께 비교하므로 참고용으로 보세요.").font = F_NOTE
ws.freeze_panes = "C6"
ws.auto_filter.ref = f"A5:Q{MENU_LAST}"

# ================================================================ 6. 매출일보
ws = wb.create_sheet("매출일보")
title(ws, "매출일보", "매일 POS 마감 후 지점별 1줄. 금액은 VAT 포함 실매출(할인 후). [주류·음료 매출]은 POS 분류별 매출에서 확인하세요.")
heads = ["날짜", "지점", "카드", "현금", "계좌이체\n·기타", "배달앱", "총매출\n(VAT포함)", "그중 주류·\n음료 매출", "음식 매출", "객수", "테이블수", "객단가", "비고"]
header_row(ws, 3, heads, height=40)
widths(ws, dict(zip("ABCDEFGHIJKLM", [12, 10, 12, 11, 11, 11, 13, 12, 13, 8, 8, 10, 24])))
for r in range(SALES_FIRST, SALES_LAST + 1):
    for col in (1, 2, 3, 4, 5, 6, 8, 10, 11, 13):
        style(ws.cell(r, col), F_INPUT, DATE if col == 1 else (NUM if col not in (2, 13) else None), FILL_INPUT, CENTER if col in (1, 2, 10, 11) else None)
    style(ws.cell(r, 7, f'=IF(COUNT(C{r}:F{r})=0,"",SUM(C{r}:F{r}))'), F_BASE, NUM)
    style(ws.cell(r, 9, f'=IF(G{r}="","",G{r}-N(H{r}))'), F_BASE, NUM)
    style(ws.cell(r, 12, f'=IF(OR(G{r}="",J{r}="",J{r}=0),"",G{r}/J{r})'), F_BASE, NUM)
import datetime as _dt
for col, v in zip(range(1, 14), [_dt.date(2026, 10, 1), "판교점", 1850000, 120000, 80000, 250000, None, 520000, None, 62, 19, None, "예시(삭제)"]):
    if v is not None:
        ws.cell(SALES_FIRST, col).value = v
add_list_dv(ws, STORE_LIST, f"B{SALES_FIRST}:B{SALES_LAST}")
ws.conditional_formatting.add(f"B{SALES_FIRST}:B{SALES_LAST}", FormulaRule(formula=[f'AND($A{SALES_FIRST}<>"",$B{SALES_FIRST}="")'], fill=RED_FILL))
ws.freeze_panes = "C4"
ws.auto_filter.ref = f"A3:M{SALES_LAST}"

# ================================================================ 7. 매입장
ws = wb.create_sheet("매입장")
title(ws, "매입장 (식자재·주류·소모품 매입)", "거래명세서·영수증 1품목당 1줄. 품목을 고르면 분류·단위·과세구분이 자동 입력(직접 수정 가능). 목록에 없는 품목은 직접 입력하고 분류를 꼭 고르세요. 수량·단가 없이 금액만 있으면 [공급가액]에 직접 입력해도 됩니다.")
heads = ["날짜", "지점", "거래처", "품목명", "분류", "수량", "단위", "단가\n(공급가)", "공급가액", "과세\n구분", "부가세", "합계\n(결제액)", "결제방법", "증빙", "비고", "원가그룹\n(자동)"]
header_row(ws, 3, heads, height=40)
widths(ws, dict(zip("ABCDEFGHIJKLMNOP", [12, 10, 14, 24, 17, 8, 7, 11, 12, 7, 10, 12, 10, 13, 18, 11])))
for r in range(PUR_FIRST, PUR_LAST + 1):
    for col in (1, 2, 3, 4, 6, 8, 13, 14, 15):
        style(ws.cell(r, col), F_INPUT, {1: DATE, 6: NUM1, 8: NUM}.get(col), FILL_INPUT, CENTER if col in (1, 2, 6, 13) else None)
    style(ws.cell(r, 5, f'=IF(D{r}="","",IFERROR({item_lookup("B", f"D{r}")},""))'), F_LINK, fill=FILL_INPUT)
    style(ws.cell(r, 7, f'=IF(D{r}="","",IFERROR({item_lookup("F", f"D{r}")},""))'), F_LINK, align=CENTER)
    style(ws.cell(r, 9, f'=IF(OR(F{r}="",H{r}=""),"",F{r}*H{r})'), F_BASE, NUM, FILL_INPUT)
    style(ws.cell(r, 10, f'=IF(D{r}="","",IFERROR({item_lookup("M", f"D{r}")},"과세"))'), F_LINK, fill=FILL_INPUT, align=CENTER)
    style(ws.cell(r, 11, f'=IF(I{r}="","",IF(J{r}="과세",ROUND(I{r}*0.1,0),0))'), F_BASE, NUM)
    style(ws.cell(r, 12, f'=IF(I{r}="","",I{r}+K{r})'), F_BASE, NUM)
    style(ws.cell(r, 16, f'=IF(E{r}="","",IFERROR(INDEX({GRP_RNG},MATCH(E{r},{CAT_RNG},0)),""))'), F_LINK, align=CENTER)
examples = [
    (_dt.date(2026, 10, 1), "판교점", "양고기 거래처", "양다리(뼈포함·냉동)", 20, 14000, "카드", "계산서(면세)"),
    (_dt.date(2026, 10, 1), "판교점", "양꼬치 거래처", "생양꼬치(완제품)", 3, 45000, "계좌이체", "계산서(면세)"),
    (_dt.date(2026, 10, 2), "판교점", "주류 도매상", "연태구냥 500ml", 24, 12000, "계좌이체", "세금계산서"),
]
for i, (d, st, vd, it, q, p, pay, ev) in enumerate(examples):
    r = PUR_FIRST + i
    for col, v in ((1, d), (2, st), (3, vd), (4, it), (6, q), (8, p), (13, pay), (14, ev), (15, "예시(삭제)")):
        ws.cell(r, col).value = v
add_list_dv(ws, STORE_LIST, f"B{PUR_FIRST}:B{PUR_LAST}")
add_list_dv(ws, f"={rng('식자재단가', 'C', ITEM_FIRST, ITEM_LAST)}", f"D{PUR_FIRST}:D{PUR_LAST}", strict=False)
add_list_dv(ws, f"={CAT_RNG}", f"E{PUR_FIRST}:E{PUR_LAST}")
add_list_dv(ws, f"={S}!$N$4:$N$5", f"J{PUR_FIRST}:J{PUR_LAST}")
add_list_dv(ws, f"={S}!$O$4:$O$7", f"M{PUR_FIRST}:M{PUR_LAST}")
add_list_dv(ws, f"={S}!$P$4:$P$8", f"N{PUR_FIRST}:N{PUR_LAST}")
ws.conditional_formatting.add(f"E{PUR_FIRST}:E{PUR_LAST}", FormulaRule(formula=[f'AND($D{PUR_FIRST}<>"",$E{PUR_FIRST}="")'], fill=RED_FILL))
ws.conditional_formatting.add(f"B{PUR_FIRST}:B{PUR_LAST}", FormulaRule(formula=[f'AND($A{PUR_FIRST}<>"",$B{PUR_FIRST}="")'], fill=RED_FILL))
ws.freeze_panes = "E4"
ws.auto_filter.ref = f"A3:P{PUR_LAST}"

# ================================================================ 8. 재고실사
ws = wb.create_sheet("재고실사")
title(ws, "월말 재고실사", "매월 말일 영업 종료 후 남은 수량을 해당 월 칸에 입력(구매단위 기준, 소수 가능: 양다리 3.5kg, 양꼬치 0.4박스). [전년말 재고]는 1월 기초재고입니다. 단가는 식자재단가에서 자동으로 불러옵니다.", span="A1:W1")
months = [f"{m}월말" for m in range(1, 13)]
heads = ["지점", "분류(자동)", "품목명", "구매\n단위", "단가\n(공급가)", "전년말\n재고"] + months + ["조회월\n기말수량", "조회월\n기말금액", "조회월\n기초수량", "조회월\n기초금액", "원가그룹\n(자동)"]
header_row(ws, 3, heads, height=40)
widths(ws, {"A": 9, "B": 16, "C": 24, "D": 7, "E": 10, "F": 8, **{get_column_letter(c): 7 for c in range(7, 19)}, "S": 9, "T": 11, "U": 9, "V": 11, "W": 10})
inv_rows = [(si, it[2]) for si in range(len(STORES)) for it in ITEMS]
for r in range(INV_FIRST, INV_LAST + 1):
    pre = inv_rows[r - INV_FIRST] if r - INV_FIRST < len(inv_rows) else None
    style(ws.cell(r, 1, f"={S}!$E${5 + pre[0]}" if pre else None), F_LINK if pre else F_INPUT, fill=FILL_INPUT, align=CENTER)
    style(ws.cell(r, 2, f'=IF(C{r}="","",IFERROR({item_lookup("B", f"C{r}")},""))'), F_LINK, fill=FILL_INPUT)
    style(ws.cell(r, 3, pre[1] if pre else None), F_INPUT, fill=FILL_INPUT)
    style(ws.cell(r, 4, f'=IF(C{r}="","",IFERROR({item_lookup("F", f"C{r}")},""))'), F_LINK, align=CENTER)
    style(ws.cell(r, 5, f'=IF(C{r}="",0,IFERROR({item_lookup("G", f"C{r}")},0))'), F_LINK, NUM, FILL_INPUT)
    for col in range(6, 19):
        style(ws.cell(r, col), F_INPUT, NUM1, FILL_INPUT, CENTER)
    style(ws.cell(r, 19, f"=N(INDEX(F{r}:R{r},1,{MONTH}+1))"), F_BASE, NUM1, align=CENTER)
    style(ws.cell(r, 20, f"=S{r}*E{r}"), F_BASE, NUM)
    style(ws.cell(r, 21, f"=N(INDEX(F{r}:R{r},1,{MONTH}))"), F_BASE, NUM1, align=CENTER)
    style(ws.cell(r, 22, f"=U{r}*E{r}"), F_BASE, NUM)
    style(ws.cell(r, 23, f'=IF(B{r}="","",IFERROR(INDEX({GRP_RNG},MATCH(B{r},{CAT_RNG},0)),""))'), F_LINK, align=CENTER)
add_list_dv(ws, STORE_LIST, f"A{INV_FIRST}:A{INV_LAST}")
add_list_dv(ws, f"={rng('식자재단가', 'C', ITEM_FIRST, ITEM_LAST)}", f"C{INV_FIRST}:C{INV_LAST}", strict=False)
add_list_dv(ws, f"={CAT_RNG}", f"B{INV_FIRST}:B{INV_LAST}")
# 조회월 열 강조
ws.conditional_formatting.add(f"G3:R{INV_LAST}", FormulaRule(formula=[f"COLUMN(G3)-6={MONTH}"], fill=PatternFill("solid", fgColor="FFE08A")))
ws.freeze_panes = "D4"
ws.auto_filter.ref = f"A3:W{INV_LAST}"

# ================================================================ 9. 매출원가계산서
ws = wb.create_sheet("매출원가계산서")
widths(ws, {"A": 2, "B": 24, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 12, "I": 12, "J": 3, "K": 10, "L": 12})
ws["B1"] = f"={S}!$C$3&\" 매출원가계산서\""
ws["B1"].font = F_TITLE
ws.row_dimensions[1].height = 30
ws["K2"], ws["K3"] = "시작일", "종료일(미만)"
ws["L2"], ws["L3"] = f"=DATE({YEAR},{MONTH},1)", f"=DATE({YEAR},{MONTH}+1,1)"
for a in ("K2", "K3", "L2", "L3"):
    ws[a].font = F_NOTE
    ws[a].number_format = DATE
ws["B2"] = f'="기간: "&TEXT($L$2,"yyyy-mm-dd")&" ~ "&TEXT($L$3-1,"yyyy-mm-dd")&"     지점: "&{STORE}&"     (금액 단위: 원, 원가는 부가세 제외 공급가액)"'
ws["B2"].font = Font(name=FONT, size=10, bold=True, color="555055")
DATES = "{d},\">=\"&$L$2,{d},\"<\"&$L$3"


def sales_sum(col):
    d = rng("매출일보", "A", SALES_FIRST, SALES_LAST)
    args = f"{rng('매출일보', col, SALES_FIRST, SALES_LAST)},{DATES.format(d=d)}"
    return by_store(args, rng("매출일보", "B", SALES_FIRST, SALES_LAST))


def section(row, text):
    ws.cell(row, 2, text).font = F_SECTION
    ws.row_dimensions[row].height = 22


# 1. 매출
section(4, "1. 매출 현황")
header_row(ws, 5, ["구분", "매출 (VAT포함)", "매출 (공급가액)", "구성비"], col=2, height=24)
for r, label, col in ((6, "음식 매출", "I"), (7, "주류·음료 매출", "H")):
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, f"={sales_sum(col)}"), F_LINK, NUM)
    style(ws.cell(r, 4, f"=C{r}/1.1"), F_BASE, NUM)
    style(ws.cell(r, 5, f"=IFERROR(D{r}/$D$8,0)"), F_BASE, PCT, align=CENTER)
style(ws["B8"], F_BOLD, fill=FILL_TOTAL).value = "매출 합계"
for col, f, fmt in (("C", "=C6+C7", NUM), ("D", "=D6+D7", NUM), ("E", "=IFERROR(D8/$D$8,0)", PCT)):
    style(ws[f"{col}8"], F_BOLD, fmt, FILL_TOTAL).value = f
style(ws["B9"], F_BOLD, fill=FILL_SUB).value = "객수 (명)"
style(ws["C9"], F_LINK, "#,##0").value = f"={sales_sum('J')}"
style(ws["B10"], F_BOLD, fill=FILL_SUB).value = "객단가 (VAT포함)"
style(ws["C10"], F_BASE, NUM).value = "=IFERROR(C8/C9,0)"

# 2. 매출원가
section(12, "2. 매출원가  =  기초재고 + 당월매입 − 기말재고")
header_row(ws, 13, ["분류", "원가그룹", "기초재고", "당월매입", "기말재고", "매출원가", "원가 구성비", "관련매출 대비"], col=2, height=28)
C_FIRST = 14
C_LAST = C_FIRST + NCAT - 1
INV_A, INV_B = rng("재고실사", "A", INV_FIRST, INV_LAST), rng("재고실사", "B", INV_FIRST, INV_LAST)
PUR_D = rng("매입장", "A", PUR_FIRST, PUR_LAST)
for i in range(NCAT):
    r = C_FIRST + i
    style(ws.cell(r, 2, f"={S}!$G${CAT_FIRST + i}"), Font(name=FONT, size=10, bold=True, color="008000"), fill=FILL_SUB)
    style(ws.cell(r, 3, f"={S}!$H${CAT_FIRST + i}"), F_LINK, align=CENTER)
    style(ws.cell(r, 4, "=" + by_store(f"{rng('재고실사', 'V', INV_FIRST, INV_LAST)},{INV_B},$B{r}", INV_A)), F_LINK, NUM)
    pur_args = f"{rng('매입장', 'I', PUR_FIRST, PUR_LAST)},{rng('매입장', 'E', PUR_FIRST, PUR_LAST)},$B{r},{DATES.format(d=PUR_D)}"
    style(ws.cell(r, 5, "=" + by_store(pur_args, rng("매입장", "B", PUR_FIRST, PUR_LAST))), F_LINK, NUM)
    style(ws.cell(r, 6, "=" + by_store(f"{rng('재고실사', 'T', INV_FIRST, INV_LAST)},{INV_B},$B{r}", INV_A)), F_LINK, NUM)
    style(ws.cell(r, 7, f"=D{r}+E{r}-F{r}"), F_BOLD, NUM)
    style(ws.cell(r, 8, f"=IFERROR(G{r}/$G${C_LAST + 4},0)"), F_BASE, PCT, align=CENTER)
    style(ws.cell(r, 9, f'=IFERROR(G{r}/IF(C{r}="식재료",$D$6,IF(C{r}="주류·음료",$D$7,$D$8)),0)'), F_BASE, PCT, align=CENTER)
SUB = {}
for j, grp in enumerate(["식재료", "주류·음료", "소모품"]):
    r = C_LAST + 1 + j
    SUB[grp] = r
    style(ws.cell(r, 2, f"{grp} 소계"), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, grp), F_BOLD, fill=FILL_SUB, align=CENTER)
    for col in "DEFG":
        style(ws[f"{col}{r}"], F_BOLD, NUM, FILL_SUB).value = f'=SUMIF($C${C_FIRST}:$C${C_LAST},$C{r},{col}${C_FIRST}:{col}${C_LAST})'
    style(ws.cell(r, 8, f"=IFERROR(G{r}/$G${C_LAST + 4},0)"), F_BOLD, PCT, FILL_SUB, CENTER)
    rel = {"식재료": "$D$6", "주류·음료": "$D$7", "소모품": "$D$8"}[grp]
    style(ws.cell(r, 9, f"=IFERROR(G{r}/{rel},0)"), F_BOLD, PCT, FILL_SUB, CENTER)
TOT = C_LAST + 4
style(ws.cell(TOT, 2, "매출원가 합계"), F_BOLD, fill=FILL_TOTAL)
style(ws.cell(TOT, 3), F_BOLD, fill=FILL_TOTAL)
for col in "DEFG":
    style(ws[f"{col}{TOT}"], F_BOLD, NUM, FILL_TOTAL).value = f"=SUM({col}{SUB['식재료']}:{col}{SUB['소모품']})"
style(ws.cell(TOT, 8, f"=IFERROR(G{TOT}/$G${TOT},0)"), F_BOLD, PCT, FILL_TOTAL, CENTER)
style(ws.cell(TOT, 9, f"=IFERROR(G{TOT}/$D$8,0)"), F_BOLD, PCT, FILL_TOTAL, CENTER)
ws.cell(TOT + 1, 2, "※ 기초재고 = 재고실사의 전월말 수량 × 현재 단가,  기말재고 = 조회월말 수량 × 현재 단가.  '관련매출 대비'는 식재료→음식매출, 주류·음료→주류매출, 소모품→총매출 기준.").font = F_NOTE

# 3. 원가율 분석
R3 = TOT + 3
section(R3, "3. 원가율 분석 (목표 대비)")
header_row(ws, R3 + 1, ["구분", "매출원가", "매출(공급가)", "실제 원가율", "목표 원가율", "차이(%p)", "판정"], col=2, height=24)
fs, ds, ss = SUB["식재료"], SUB["주류·음료"], SUB["소모품"]
rows3 = [
    ("식재료 원가율", f"=G{fs}", "=D6", f"={TGT_FOOD}"),
    ("주류·음료 원가율", f"=G{ds}", "=D7", f"={TGT_DRINK}"),
    ("식재료+주류 원가율", f"=G{fs}+G{ds}", "=D8", f"=IFERROR(({TGT_FOOD}*D6+{TGT_DRINK}*D7)/D8,{TGT_FOOD})"),
    ("소모품 포함 총원가율", f"=G{TOT}", "=D8", None),
]
for i, (label, cost, sale, tgt) in enumerate(rows3):
    r = R3 + 2 + i
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, cost), F_BASE, NUM)
    style(ws.cell(r, 4, sale), F_BASE, NUM)
    style(ws.cell(r, 5, f"=IFERROR(C{r}/D{r},0)"), F_BOLD, PCT, align=CENTER)
    style(ws.cell(r, 6, tgt), F_LINK if tgt else F_BASE, PCT, align=CENTER)
    style(ws.cell(r, 7, f'=IF(F{r}="","",E{r}-F{r})' if tgt else None), F_BASE, PCT, align=CENTER)
    judge = f'=IF(D{r}=0,"매출 없음",IF(E{r}>F{r}+{TOL},"위험",IF(E{r}>F{r},"주의","적정")))' if tgt else "참고용"
    style(ws.cell(r, 8, judge), F_BOLD, align=CENTER)
status_cf(ws, f"H{R3 + 2}:H{R3 + 5}")

# 4. 매출총이익
R4 = R3 + 7
section(R4, "4. 매출총이익")
rows4 = [("매출 (공급가액)", "=D8", NUM), ("(−) 매출원가", f"=G{TOT}", NUM), ("매출총이익", f"=C{R4 + 1}-C{R4 + 2}", NUM),
         ("매출총이익률", f"=IFERROR(C{R4 + 3}/C{R4 + 1},0)", PCT), ("1인당 매출원가 (객수 기준)", f"=IFERROR(G{TOT}/C9,0)", NUM)]
for i, (label, f, fmt) in enumerate(rows4):
    r = R4 + 1 + i
    bold = label.startswith("매출총이익")
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_TOTAL if bold else FILL_SUB)
    style(ws.cell(r, 3, f), F_BOLD if bold else F_BASE, fmt, FILL_TOTAL if bold else None)
ws.cell(R4 + 6, 2, "※ 인건비·임대료·공과금 등 판매관리비는 포함하지 않은 '매출총이익'입니다.").font = F_NOTE

# 5. 이론 vs 실제
R5 = R4 + 8
section(R5, "5. 이론원가(레시피) vs 실제원가 — 로스 점검")
header_row(ws, R5 + 1, ["구분", "이론원가", "실제원가", "차이(실제−이론)", "차이율", "점검 결과"], col=2, height=24)
M_C, M_O = rng("메뉴원가표", "C", MENU_FIRST, MENU_LAST), rng("메뉴원가표", "O", MENU_FIRST, MENU_LAST)
for i, (label, grp, actual) in enumerate([("식재료", "식재료", f"=G{fs}"), ("주류·음료", "주류·음료", f"=G{ds}")]):
    r = R5 + 2 + i
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, f'=SUMIF({M_C},"{grp}",{M_O})'), F_LINK, NUM)
    style(ws.cell(r, 4, actual), F_BASE, NUM)
    style(ws.cell(r, 5, f"=D{r}-C{r}"), F_BASE, NUM)
    style(ws.cell(r, 6, f"=IFERROR(E{r}/C{r},0)"), F_BASE, PCT, align=CENTER)
    style(ws.cell(r, 7, f'=IF(C{r}=0,"메뉴원가표에 판매수량 입력 필요",IF(F{r}>{GAP},"주의: 로스·과다투입·서비스·누락 점검",IF(F{r}<-{GAP},"확인: 매입 누락·재고 과다계상 의심","적정 범위")))'), F_BOLD, align=LEFT)
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
ws.conditional_formatting.add(f"G{R5 + 2}:G{R5 + 3}", FormulaRule(formula=[f'LEFT(G{R5 + 2},2)="주의"'], fill=RED_FILL, font=Font(name=FONT, color="C8322A", bold=True)))
ws.conditional_formatting.add(f"G{R5 + 2}:G{R5 + 3}", FormulaRule(formula=[f'LEFT(G{R5 + 2},2)="확인"'], fill=ORANGE_FILL))
ws.conditional_formatting.add(f"G{R5 + 2}:G{R5 + 3}", FormulaRule(formula=[f'LEFT(G{R5 + 2},2)="적정"'], fill=GREEN_FILL, font=Font(name=FONT, color="187A56", bold=True)))
ws.cell(R5 + 4, 2, "※ 이론원가는 전 지점 합산 판매수량 기준입니다. 지점별로 보려면 해당 지점 판매수량을 메뉴원가표에 입력하세요.").font = F_NOTE

# 6. 의제매입
R6 = R5 + 6
section(R6, "6. 참고: 의제매입세액 공제 예상 (부가세 신고용)")
tax_args = f"{rng('매입장', 'I', PUR_FIRST, PUR_LAST)},{rng('매입장', 'J', PUR_FIRST, PUR_LAST)},\"면세\",{rng('매입장', 'P', PUR_FIRST, PUR_LAST)},\"식재료\",{DATES.format(d=PUR_D)}"
rows6 = [("당월 면세 식재료 매입액", "=" + by_store(tax_args, rng("매입장", "B", PUR_FIRST, PUR_LAST)), NUM, F_LINK),
         ("공제율", f"={DEDUCT}", "0.00%", F_LINK),
         ("공제 예상액", f"=ROUND(C{R6 + 1}*C{R6 + 2},0)", NUM, F_BOLD)]
for i, (label, f, fmt, font) in enumerate(rows6):
    r = R6 + 1 + i
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, f), font, fmt, FILL_TOTAL if i == 2 else None)
ws.cell(R6 + 4, 2, "※ 계산서·신용카드·현금영수증 등 적격증빙이 있는 면세 농축수산물만 공제됩니다. 과세표준 대비 한도가 있으니 신고 시 세무사와 확인하세요.").font = F_NOTE

# 7. 점검
R7 = R6 + 6
section(R7, "7. 입력 점검")
checks = [
    ("조회월 기말재고 입력", f'=IF(SUMPRODUCT(--(INDEX(\'재고실사\'!$G${INV_FIRST}:$R${INV_LAST},0,{MONTH})<>""))=0,"미입력 → 기말재고 0으로 계산되어 원가가 과대 표시됩니다","입력됨")'),
    ("매입장 분류 누락 건수", f'=COUNTIFS({rng("매입장", "D", PUR_FIRST, PUR_LAST)},"<>",{rng("매입장", "E", PUR_FIRST, PUR_LAST)},"")&" 건"'),
    ("매입장 지점 누락 건수", f'=COUNTIFS({PUR_D},"<>",{rng("매입장", "B", PUR_FIRST, PUR_LAST)},"")&" 건"'),
    ("레시피 미입력 메뉴 수", f'=COUNTIF({rng("메뉴원가표", "K", MENU_FIRST, MENU_LAST)},"레시피 미입력")&" 개"'),
    ("원가율 '위험' 메뉴 수", f'=COUNTIF({rng("메뉴원가표", "K", MENU_FIRST, MENU_LAST)},"위험")&" 개"'),
]
for i, (label, f) in enumerate(checks):
    r = R7 + 1 + i
    style(ws.cell(r, 2, label), F_BOLD, fill=FILL_SUB)
    style(ws.cell(r, 3, f), F_BASE, align=LEFT)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
ws.conditional_formatting.add(f"C{R7 + 1}", FormulaRule(formula=[f'LEFT(C{R7 + 1},3)="미입력"'], fill=RED_FILL, font=Font(name=FONT, color="C8322A", bold=True)))
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"B1:I{R7 + 6}"
ws.sheet_view.showGridLines = False

# ================================================================ 10. 연간요약
ws = wb.create_sheet("연간요약")
title(ws, "연간 요약 (월별 원가율 추이)", "기본설정의 연도·조회지점 기준. 재고실사를 하지 않은 달은 원가가 부정확하니 [재고실사] 열을 확인하세요.", span="A1:S1")
heads = ["월", "총매출\n(VAT포함)", "음식매출\n(공급가)", "주류매출\n(공급가)", "매출 합계\n(공급가)", "식재료\n매입", "식재료\n원가", "식재료\n원가율", "주류·음료\n매입", "주류·음료\n원가", "주류\n원가율", "소모품\n원가", "총\n매출원가", "총\n원가율", "매출\n총이익", "총이익률", "객수", "객단가", "재고실사"]
header_row(ws, 4, heads, height=40)
widths(ws, {get_column_letter(c): w for c, w in zip(range(1, 20), [6, 13, 12, 12, 13, 12, 12, 8, 11, 11, 8, 10, 12, 8, 12, 8, 8, 9, 9])})
Y_FIRST, Y_LAST = 5, 16
INV_E, INV_W = rng("재고실사", "E", INV_FIRST, INV_LAST), rng("재고실사", "W", INV_FIRST, INV_LAST)
INV_QTY = f"'재고실사'!$F${INV_FIRST}:$R${INV_LAST}"


def y_dates(r):
    d = rng("매출일보", "A", SALES_FIRST, SALES_LAST)
    return d, f'{d},">="&DATE({YEAR},$A{r},1),{d},"<"&DATE({YEAR},$A{r}+1,1)'


def y_sales(col, r):
    _, cond = y_dates(r)
    return by_store(f"{rng('매출일보', col, SALES_FIRST, SALES_LAST)},{cond}", rng("매출일보", "B", SALES_FIRST, SALES_LAST))


def y_purchase(grp, r):
    cond = f'{PUR_D},">="&DATE({YEAR},$A{r},1),{PUR_D},"<"&DATE({YEAR},$A{r}+1,1)'
    return by_store(f"{rng('매입장', 'I', PUR_FIRST, PUR_LAST)},{rng('매입장', 'P', PUR_FIRST, PUR_LAST)},\"{grp}\",{cond}", rng("매입장", "B", PUR_FIRST, PUR_LAST))


def y_inv(grp, r, offset):
    q = f"INDEX({INV_QTY},0,$A{r}+{offset})"
    return (f'IF({STORE}="전체",SUMPRODUCT(({INV_W}="{grp}")*{INV_E}*{q}),'
            f'SUMPRODUCT(({INV_W}="{grp}")*({INV_A}={STORE})*{INV_E}*{q}))')


def y_cost(grp, r):
    return f"{y_inv(grp, r, 0)}+{y_purchase(grp, r)}-{y_inv(grp, r, 1)}"


for r in range(Y_FIRST, Y_LAST + 1):
    m = r - Y_FIRST + 1
    f = {
        1: (m, '0"월"'),
        2: (f"={y_sales('G', r)}", NUM),
        3: (f"={y_sales('I', r)}/1.1", NUM),
        4: (f"={y_sales('H', r)}/1.1", NUM),
        5: (f"=C{r}+D{r}", NUM),
        6: (f"={y_purchase('식재료', r)}", NUM),
        7: (f"={y_cost('식재료', r)}", NUM),
        8: (f"=IFERROR(G{r}/C{r},0)", PCT),
        9: (f"={y_purchase('주류·음료', r)}", NUM),
        10: (f"={y_cost('주류·음료', r)}", NUM),
        11: (f"=IFERROR(J{r}/D{r},0)", PCT),
        12: (f"={y_cost('소모품', r)}", NUM),
        13: (f"=G{r}+J{r}+L{r}", NUM),
        14: (f"=IFERROR(M{r}/E{r},0)", PCT),
        15: (f"=E{r}-M{r}", NUM),
        16: (f"=IFERROR(O{r}/E{r},0)", PCT),
        17: (f"={y_sales('J', r)}", "#,##0"),
        18: (f"=IFERROR(B{r}/Q{r},0)", NUM),
        19: (f'=IF(SUMPRODUCT(--(INDEX({INV_QTY},0,$A{r}+1)<>""))>0,"입력","미입력")', None),
    }
    for col, (v, fmt) in f.items():
        style(ws.cell(r, col, v), F_BOLD if col == 1 else F_BASE, fmt, FILL_SUB if col == 1 else None, CENTER if col in (1, 8, 11, 14, 16, 19) else None)
T = Y_LAST + 1
style(ws.cell(T, 1, "합계"), F_BOLD, fill=FILL_TOTAL, align=CENTER)
for col in range(2, 20):
    L = get_column_letter(col)
    v = {8: f"=IFERROR(G{T}/C{T},0)", 11: f"=IFERROR(J{T}/D{T},0)", 14: f"=IFERROR(M{T}/E{T},0)", 16: f"=IFERROR(O{T}/E{T},0)",
         18: f"=IFERROR(B{T}/Q{T},0)", 19: None}.get(col, f"=SUM({L}{Y_FIRST}:{L}{Y_LAST})")
    style(ws.cell(T, col, v), F_BOLD, PCT if col in (8, 11, 14, 16) else ("#,##0" if col == 17 else NUM), FILL_TOTAL, CENTER if col in (8, 11, 14, 16) else None)
for col, tgt in (("H", TGT_FOOD), ("K", TGT_DRINK)):
    ws.conditional_formatting.add(f"{col}{Y_FIRST}:{col}{Y_LAST}", FormulaRule(formula=[f"AND({col}{Y_FIRST}>0,{col}{Y_FIRST}>{tgt}+{TOL})"], fill=RED_FILL, font=Font(name=FONT, color="C8322A", bold=True)))
    ws.conditional_formatting.add(f"{col}{Y_FIRST}:{col}{Y_LAST}", FormulaRule(formula=[f"AND({col}{Y_FIRST}>0,{col}{Y_FIRST}>{tgt})"], fill=ORANGE_FILL))
    ws.conditional_formatting.add(f"{col}{Y_FIRST}:{col}{Y_LAST}", FormulaRule(formula=[f"AND({col}{Y_FIRST}>0,{col}{Y_FIRST}<={tgt})"], fill=GREEN_FILL))
ws.conditional_formatting.add(f"S{Y_FIRST}:S{Y_LAST}", CellIsRule(operator="equal", formula=['"미입력"'], font=Font(name=FONT, color="C8322A")))
ws.cell(T + 2, 1, "※ 원가율 색상: 초록 = 목표 이하, 주황 = 목표 초과, 빨강 = 목표 + 허용폭 초과.").font = F_NOTE
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- 마무리
for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 100
for name, color in {"사용법": DARK, "기본설정": DARK, "매출원가계산서": CORAL, "연간요약": CORAL}.items():
    wb[name].sheet_properties.tabColor = color
wb.move_sheet("매출원가계산서", offset=-(wb.sheetnames.index("매출원가계산서") - 2))
wb.move_sheet("연간요약", offset=-(wb.sheetnames.index("연간요약") - 3))
wb.active = wb.sheetnames.index("매출원가계산서")
for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = sheet.title == "매출원가계산서"
wb.save(OUT)
print("saved", OUT)
